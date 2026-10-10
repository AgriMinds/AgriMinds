"""Comprehensive Publication & Project Documentation Generator for AgriMinds (AI-DREWS).
Generates a 30+ page equivalent, research-grade, publication-ready DOCX document
containing all architectural specifications, mathematical formulations, comparative evaluations,
historical validation benchmarks, agro-ecological rules, database designs, API contracts,
and visual figures/screenshots.
"""

import os
import sys
import datetime
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    """Set shading color for a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Set internal cell margins (padding) in dxa (1 pt = 20 dxa)."""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(
        f'<w:tcMar {nsdecls("w")}>'
        f'<w:top w:w="{top}" w:type="dxa"/>'
        f'<w:bottom w:w="{bottom}" w:type="dxa"/>'
        f'<w:left w:w="{left}" w:type="dxa"/>'
        f'<w:right w:w="{right}" w:type="dxa"/>'
        f'</w:tcMar>'
    )
    tcPr.append(tcMar)

def set_cell_borders(cell, top=None, bottom=None, left=None, right=None):
    """Set custom borders on a table cell."""
    tcPr = cell._tc.get_or_add_tcPr()
    borders_elm = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    
    borders = {'top': top, 'bottom': bottom, 'left': left, 'right': right}
    for side, border_style in borders.items():
        if border_style:
            # border_style: (val, sz, color)
            val, sz, col = border_style
            b_elm = parse_xml(f'<w:{side} {nsdecls("w")} w:val="{val}" w:sz="{sz}" w:space="0" w:color="{col}"/>')
            borders_elm.append(b_elm)
        else:
            b_elm = parse_xml(f'<w:{side} {nsdecls("w")} w:val="none"/>')
            borders_elm.append(b_elm)
    tcPr.append(borders_elm)

def format_table(table, col_widths, header_bg="1B5E20", alt_bg="F1F8E9"):
    """Format table with column widths, borders, repeating headers, and zebra rows."""
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    # Prevent row split across pages and repeat header
    for i, row in enumerate(table.rows):
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if i == 0:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))

        for j, cell in enumerate(row.cells):
            cell.width = col_widths[j]
            set_cell_margins(cell, top=120, bottom=120, left=160, right=160)
            if i == 0:
                set_cell_background(cell, header_bg)
                set_cell_borders(cell,
                                 top=('single', '8', '1B5E20'),
                                 bottom=('single', '16', '0A3912'),
                                 left=('none', '0', 'auto'),
                                 right=('none', '0', 'auto'))
                for p in cell.paragraphs:
                    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                    for run in p.runs:
                        run.font.bold = True
                        run.font.color.rgb = RGBColor(255, 255, 255)
                        run.font.size = Pt(9.5)
            else:
                bg = alt_bg if (i % 2 == 1) else "FFFFFF"
                set_cell_background(cell, bg)
                set_cell_borders(cell,
                                 top=('single', '4', 'E0E0E0'),
                                 bottom=('single', '4', 'E0E0E0'),
                                 left=('none', '0', 'auto'),
                                 right=('none', '0', 'auto'))
                for p in cell.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9.0)
                        run.font.color.rgb = RGBColor(33, 33, 33)

def add_callout(doc, text_list, title=None, alert_type="NOTE"):
    """Add a professional GitHub-style alert / callout box."""
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    table.rows[0].cells[0].width = Inches(6.5)
    
    cell = table.rows[0].cells[0]
    set_cell_margins(cell, top=160, bottom=160, left=240, right=200)
    
    config = {
        "NOTE": {"bg": "E8F5E9", "border": "2E7D32", "title_col": RGBColor(46, 125, 50), "prefix": "NOTE: "},
        "IMPORTANT": {"bg": "E3F2FD", "border": "1565C0", "title_col": RGBColor(21, 101, 192), "prefix": "IMPORTANT: "},
        "WARNING": {"bg": "FFF3E0", "border": "E65100", "title_col": RGBColor(230, 81, 0), "prefix": "WARNING: "},
        "TIP": {"bg": "EDE7F6", "border": "512DA8", "title_col": RGBColor(81, 45, 168), "prefix": "TIP: "}
    }
    cfg = config.get(alert_type, config["NOTE"])
    
    set_cell_background(cell, cfg["bg"])
    set_cell_borders(cell,
                     left=('single', '36', cfg["border"]),
                     top=('none', '0', 'auto'),
                     bottom=('none', '0', 'auto'),
                     right=('none', '0', 'auto'))
    
    p = cell.paragraphs[0]
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    r_prefix = p.add_run(cfg["prefix"])
    r_prefix.bold = True
    r_prefix.font.size = Pt(10)
    r_prefix.font.color.rgb = cfg["title_col"]
    
    if title:
        r_title = p.add_run(f"{title}\n")
        r_title.bold = True
        r_title.font.size = Pt(10)
        r_title.font.color.rgb = cfg["title_col"]
        
    for idx, t in enumerate(text_list):
        if idx > 0 or title:
            p2 = cell.add_paragraph()
            p2.paragraph_format.space_before = Pt(2)
            p2.paragraph_format.space_after = Pt(3)
            p2.paragraph_format.line_spacing = 1.15
            run = p2.add_run(t)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(33, 33, 33)
        else:
            run = p.add_run(t)
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(33, 33, 33)
            
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def add_heading_with_spacing(doc, text, level):
    h = doc.add_heading(text, level=level)
    h.paragraph_format.keep_with_next = True
    if level == 1:
        h.paragraph_format.space_before = Pt(18)
        h.paragraph_format.space_after = Pt(8)
        for r in h.runs:
            r.font.size = Pt(16)
            r.font.bold = True
            r.font.color.rgb = RGBColor(27, 94, 32) # Dark Forest Green
    elif level == 2:
        h.paragraph_format.space_before = Pt(14)
        h.paragraph_format.space_after = Pt(6)
        for r in h.runs:
            r.font.size = Pt(13)
            r.font.bold = True
            r.font.color.rgb = RGBColor(13, 71, 161) # Deep Navy
    elif level == 3:
        h.paragraph_format.space_before = Pt(10)
        h.paragraph_format.space_after = Pt(4)
        for r in h.runs:
            r.font.size = Pt(11)
            r.font.bold = True
            r.font.color.rgb = RGBColor(55, 71, 79) # Slate Charcoal
    return h

def add_caption(doc, caption_text, figure_number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(14)
    run_bold = p.add_run(f"Figure {figure_number}: " if figure_number else "Caption: ")
    run_bold.bold = True
    run_bold.font.size = Pt(9.0)
    run_bold.font.color.rgb = RGBColor(33, 33, 33)
    run_text = p.add_run(caption_text)
    run_text.italic = True
    run_text.font.size = Pt(9.0)
    run_text.font.color.rgb = RGBColor(66, 66, 66)

def add_table_caption(doc, caption_text, table_number=None):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.keep_with_next = True
    run_bold = p.add_run(f"Table {table_number}: " if table_number else "Table: ")
    run_bold.bold = True
    run_bold.font.size = Pt(9.5)
    run_bold.font.color.rgb = RGBColor(27, 94, 32)
    run_text = p.add_run(caption_text)
    run_text.italic = True
    run_text.font.size = Pt(9.5)
    run_text.font.color.rgb = RGBColor(33, 33, 33)

def add_p(doc, text, space_after=6, bold=False, italic=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = 1.15
    run = p.add_run(text)
    run.font.size = Pt(10.0)
    run.font.color.rgb = RGBColor(33, 33, 33)
    run.bold = bold
    run.italic = italic
    return p

print("Builder helper definitions loaded successfully.")
